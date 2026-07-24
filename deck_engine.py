"""Template-faithful engine: reuse the sample template's title + body text style,
only swapping content. No custom shapes (no accent bars, bands, callouts, tables,
or images)."""
import copy, math
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

FONT = "TT Hoves"
FONT_BOLD = "TT Hoves ExtraBold"

# template body box geometry (from slide "Physical Quantities")
BODY_L = 0.7645
BODY_T = 1.68
BODY_W = 11.67
BODY_BOTTOM = 6.35            # keep clear of footer (~6.51in)


# ------------------------------------------------------------------ #
def duplicate_slide(prs, src_slide):
    layout = src_slide.slide_layout
    new_slide = prs.slides.add_slide(layout)
    for shp in list(new_slide.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in src_slide.shapes:
        new_slide.shapes._spTree.append(copy.deepcopy(shp._element))
    return new_slide


def find_shape(slide, name):
    for s in slide.shapes:
        if s.name == name:
            return s
    return None


def delete_sldIds(prs, sldId_elements):
    lst = prs.slides._sldIdLst
    for el in sldId_elements:
        lst.remove(el)


# ------------------------------------------------------------------ #
#  Inline emphasis parser  ***bold+italic** / **bold** / *italic*
# ------------------------------------------------------------------ #
def parse_inline(text):
    runs, buf = [], ""
    bold = ital = False
    i, n = 0, len(text)

    def flush():
        nonlocal buf
        if buf:
            runs.append((bold, ital, buf)); buf = ""

    while i < n:
        if text[i:i+3] == "***":
            flush(); bold = not bold; ital = not ital; i += 3
        elif text[i:i+2] == "**":
            flush(); bold = not bold; i += 2
        elif text[i] == "*":
            flush(); ital = not ital; i += 1
        else:
            buf += text[i]; i += 1
    flush()
    return runs


# ------------------------------------------------------------------ #
#  Run builder (supports sub/superscript for equations)
# ------------------------------------------------------------------ #
def _baseline(run, val):
    run._r.get_or_add_rPr().set("baseline", str(val))


def add_seg(paragraph, seg, size, font=FONT):
    """seg = (kind, text). kind in t,b,i,bi,var,sub,sup,subi."""
    kind, text = seg
    r = paragraph.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    if kind in ("b", "bi"):
        r.font.bold = True
    if kind in ("i", "bi", "var", "subi"):
        r.font.italic = True
    if kind in ("sub", "subi"):
        _baseline(r, -25000); r.font.size = Pt(size * 0.92)
    if kind in ("sup",):
        _baseline(r, 30000); r.font.size = Pt(size * 0.92)
    return r


def _seg_len(seglist):
    return sum(len(s[1]) for s in seglist)


# ------------------------------------------------------------------ #
#  Title  (preserve template's orange-gradient ExtraBold run; adapt size)
# ------------------------------------------------------------------ #
def _title_size(text):
    L = len(text)
    if L <= 24:
        return 44
    if L <= 34:
        return 38
    if L <= 44:
        return 33
    if L <= 56:
        return 28
    return 25


def set_title(slide, text):
    tb = find_shape(slide, "TextBox 9")
    tf = tb.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    runs = p0.runs
    if runs:
        r = runs[0]
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
    else:
        r = p0.add_run()
    r.text = text
    r.font.size = Pt(_title_size(text))   # keeps gradFill / ExtraBold in rPr
    return tb


# ------------------------------------------------------------------ #
#  Body  (reuse template 'TextBox 8'; template bullet style, 28pt down-fit)
# ------------------------------------------------------------------ #
def _set_bullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", "457200"); pPr.set("indent", "-457200")
    buFont = pPr.makeelement(qn("a:buFont"),
                             {"typeface": "Arial", "pitchFamily": "34", "charset": "0"})
    buChar = pPr.makeelement(qn("a:buChar"), {"char": "\u2022"})
    pPr.append(buFont); pPr.append(buChar)


def _no_bullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def _fit_size(items, width, height, max_size=28, min_size=15):
    for size in range(max_size, min_size - 1, -1):
        cpl = max(8, int(width * 100.0 / size))
        lines = 0
        for it in items:
            if isinstance(it, str):
                length = len(it)
            else:                       # ('EQ'|'BSEG', seglist)
                length = _seg_len(it[1])
            lines += max(1, math.ceil(length / cpl))
        total = lines * (size * 1.25 / 72.0) + len(items) * (7.0 / 72.0)
        if total <= height:
            return size
    return min_size


def add_body(slide, items, top=BODY_T, height=None, left=BODY_L, width=BODY_W,
             size=None):
    """items: list of str (bullet) | ('EQ', segs) | ('BSEG', segs)."""
    if height is None:
        height = BODY_BOTTOM - top
    tb = find_shape(slide, "TextBox 8")
    tf = tb.text_frame
    # reposition / clear
    tb.left = Inches(left); tb.top = Inches(top)
    tb.width = Inches(width); tb.height = Inches(height)
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for p in list(tf.paragraphs):
        p._p.getparent().remove(p._p)
    if size is None:
        size = _fit_size(items, width, height)
    first = True
    for it in items:
        p = tf.add_paragraph()
        p.line_spacing = 1.04
        p.space_after = Pt(7 if size >= 20 else 5)
        p.space_before = Pt(0)
        if isinstance(it, str):
            _set_bullet(p)
            p.alignment = PP_ALIGN.LEFT
            for (bd, itc, chunk) in parse_inline(it):
                kind = "bi" if (bd and itc) else "b" if bd else "i" if itc else "t"
                add_seg(p, (kind, chunk), size)
        else:
            kind0, segs = it
            if kind0 == "EQ":
                _no_bullet(p)
                p.alignment = PP_ALIGN.CENTER
                p.space_before = Pt(4)
                for seg in segs:
                    # equations rendered bold for emphasis, same font family
                    k, txt = seg
                    bk = k
                    if k == "t":
                        bk = "b"
                    elif k == "var":
                        bk = "bi"
                    elif k == "sub":
                        bk = "sub"
                    elif k == "sup":
                        bk = "sup"
                    add_seg(p, (bk, txt), size + 2)
            else:  # BSEG bullet with inline segments
                _set_bullet(p)
                p.alignment = PP_ALIGN.LEFT
                for seg in segs:
                    add_seg(p, seg, size)
        first = False
    return tb, size


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
