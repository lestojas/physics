"""
Kinematics Guided Derivation Worksheet
- Landscape 13" x 8.5" bond paper, narrow margins
- 4 parts side by side in columns
- Header order per part: Activity / Name / Grade & Section / Date -> THEN Part title
- Dashed cut lines between columns extend the FULL height of the page (fixed row height)
- Answer space is maximized and balanced across parts based on step count
- Only light-gray shading used (no dark bands)
- Answer Key (page 2) = identical layout, with answers written in red
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

# ---------------------------------------------------------------
# Layout constants
# ---------------------------------------------------------------
PAGE_W = Inches(13)
PAGE_H = Inches(8.5)
MARGIN = Inches(0.3)
COL_WIDTH = Inches(3.1)

LIGHT_GRAY = "E7E7E7"          # the ONLY shading color used
CUT_LINE_COLOR = "999999"      # dashed cut-line color
TEXT_DARK = (40, 40, 40)
TEXT_MED = (85, 85, 85)
TEXT_LIGHT = (165, 165, 165)
TEXT_RED = (190, 0, 0)

# --- Text-height estimation constants (used to balance leftover space) ---
CHARS_PER_LINE = 50        # approx. characters that fit on one line of a 3.1in column at 7.5pt
LINE_HEIGHT_PT = 9.5        # approx. rendered height of one line of 7.5pt text
STEP_PARA_SPACING_PT = 5.0  # before(3)+after(2) spacing on each instruction paragraph

# Fixed header block height (Activity / Name / Grade&Section-Date / Part title / Goal)
# is the same for every part, so it's computed once.
HEADER_BLOCK_PT = 72.0

# Target usable content height inside each column (points). Matches the row
# height used for the table so the dashed cut-lines run the full page length.
WORKSHEET_TARGET_PT = Inches(7.85).pt - 6   # minus small cell-margin allowance
ANSWERKEY_TARGET_PT = Inches(7.55).pt - 6

MIN_ANSWER_SPACE_PT = 30.0  # floor, so no step ever gets a cramped answer area


def estimate_step_text_height(text):
    lines = max(1, -(-len(text) // CHARS_PER_LINE))  # ceil division
    return lines * LINE_HEIGHT_PT


def compute_answer_spaces(steps, target_pt):
    """Return a list of per-step 'blank writing space' point values so that
    the WHOLE column (header + instructions + writing space) fills target_pt,
    maximizing the space actually available for student responses."""
    text_heights = [estimate_step_text_height(s) for s in steps]
    total_text = sum(text_heights)
    total_para_spacing = STEP_PARA_SPACING_PT * len(steps)
    remaining = target_pt - HEADER_BLOCK_PT - total_text - total_para_spacing
    per_step = remaining / len(steps) if steps else 0
    per_step = max(MIN_ANSWER_SPACE_PT, per_step)
    return [per_step] * len(steps)


# ---------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------
def set_cell_margins(cell, top=0, bottom=0, left=70, right=70):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def add_run(paragraph, text, bold=False, italic=False, size=Pt(9), color=None, font_name='Calibri'):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = size
    run.font.name = font_name
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def set_paragraph_spacing(paragraph, before=0, after=0, line=None):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line:
        pf.line_spacing = line


def set_paragraph_shading(paragraph, color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color}"/>')
    pPr.append(shd)


def set_paragraph_bottom_border(paragraph, color="BFBFBF", sz=4):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="2" w:color="{color}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)


def set_table_borders_none(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="nil"/>'
        f'<w:left w:val="nil"/>'
        f'<w:bottom w:val="nil"/>'
        f'<w:right w:val="nil"/>'
        f'<w:insideH w:val="nil"/>'
        f'<w:insideV w:val="nil"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    # Fixed layout so column widths and row height are respected exactly
    layout = parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>')
    tblPr.append(layout)


def set_dashed_right_border(cell):
    """Dashed right border on a cell -- functions as a cut line."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:right w:val="dashed" w:sz="6" w:color="{CUT_LINE_COLOR}" w:space="0"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)


def set_row_exact_height(row, height):
    """Force a table row to an exact height so cut lines run the full length,
    even when a column's content is short."""
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trHeight = parse_xml(
        f'<w:trHeight {nsdecls("w")} w:val="{height}" w:hRule="exact"/>'
    )
    trPr.append(trHeight)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    row.height = height


# ---------------------------------------------------------------
# Column content builder
# ---------------------------------------------------------------
def build_part_content(cell, part, is_answer_key=False, target_pt=WORKSHEET_TARGET_PT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell, top=45, bottom=20, left=75, right=75)

    cell.paragraphs[0].clear()

    # --- 1. ACTIVITY NAME (first, above everything) ---
    act_para = cell.paragraphs[0]
    set_paragraph_spacing(act_para, before=0, after=1)
    add_run(act_para, "Activity: ", bold=True, size=Pt(7.5), color=TEXT_MED)
    add_run(act_para, "Guided Derivation of Kinematic Equations", size=Pt(7.5), color=(60, 60, 60))

    # --- 2. NAME ---
    name_para = cell.add_paragraph()
    set_paragraph_spacing(name_para, before=1, after=1)
    add_run(name_para, "Name: ", bold=True, size=Pt(7.5), color=TEXT_MED)
    add_run(name_para, "_" * 38, size=Pt(7.5), color=TEXT_LIGHT)

    # --- 3. GRADE & SECTION / DATE ---
    gs_para = cell.add_paragraph()
    set_paragraph_spacing(gs_para, before=1, after=4)
    add_run(gs_para, "Grade & Section: ", bold=True, size=Pt(7.5), color=TEXT_MED)
    add_run(gs_para, "_" * 13 + "   ", size=Pt(7.5), color=TEXT_LIGHT)
    add_run(gs_para, "Date: ", bold=True, size=Pt(7.5), color=TEXT_MED)
    add_run(gs_para, "_" * 11, size=Pt(7.5), color=TEXT_LIGHT)

    # --- 4. PART TITLE (comes AFTER the header block) ---
    title_para = cell.add_paragraph()
    set_paragraph_spacing(title_para, before=2, after=3)
    set_paragraph_shading(title_para, LIGHT_GRAY)
    add_run(title_para, f" PART {part['letter']}:  ", bold=True, size=Pt(9.5), color=TEXT_DARK)
    add_run(title_para, part["title"] + " ", bold=True, size=Pt(9.5), color=TEXT_DARK)

    # --- 5. GOAL LINE ---
    goal_para = cell.add_paragraph()
    set_paragraph_spacing(goal_para, before=0, after=6)
    set_paragraph_bottom_border(goal_para, color="CFCFCF", sz=4)
    add_run(goal_para, "Goal: Derive  ", size=Pt(7.5), color=TEXT_MED)
    add_run(goal_para, part["goal"], bold=True, italic=True, size=Pt(9), color=TEXT_DARK)

    # --- 6. STEPS (one algebraic action each) with maximized, balanced answer space ---
    space_values = compute_answer_spaces(part["steps"], target_pt)

    for i, step_text in enumerate(part["steps"], 1):
        step_p = cell.add_paragraph()
        set_paragraph_spacing(step_p, before=3, after=2)
        add_run(step_p, f"{i}.  ", bold=True, size=Pt(7.5), color=(60, 60, 60))
        add_run(step_p, step_text, size=Pt(7.5), color=(50, 50, 50))

        ans_p = cell.add_paragraph()
        space_pts = space_values[i - 1]
        if is_answer_key:
            set_paragraph_spacing(ans_p, before=2, after=max(6, space_pts - 12))
            add_run(ans_p, "      ")
            add_run(ans_p, part["answers"][i - 1], bold=True, italic=True, size=Pt(9), color=TEXT_RED)
        else:
            set_paragraph_spacing(ans_p, before=0, after=space_pts)
            add_run(ans_p, "", size=Pt(1))


def build_table(doc, parts, is_answer_key, row_height, target_pt):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_none(table)

    for col in table.columns:
        for cell in col.cells:
            cell.width = COL_WIDTH

    row = table.rows[0]
    for col_idx, part in enumerate(parts):
        cell = row.cells[col_idx]
        build_part_content(cell, part, is_answer_key=is_answer_key, target_pt=target_pt)
        if col_idx < 3:
            set_dashed_right_border(cell)

    set_row_exact_height(row, row_height)
    return table


# ---------------------------------------------------------------
# Content data
# ---------------------------------------------------------------
def get_parts():
    return [
        {
            "letter": "A",
            "title": "Deriving Kinematic Equation 1",
            "goal": "v = v\u2080 + at",
            "steps": [
                "Acceleration is defined as the rate of change of velocity over time. Write the defining formula for acceleration using delta notation.",
                "The change in velocity is \u0394v = v \u2212 v\u2080, and if timing starts at t = 0, then \u0394t = t. Substitute these into the definition you wrote above.",
                "Multiply both sides of your equation by t to eliminate the fraction. Write the result.",
                "Add v\u2080 to both sides of your equation to isolate v. Write the final equation.",
            ],
            "answers": [
                "a = \u0394v / \u0394t",
                "a = (v \u2212 v\u2080) / t",
                "at = v \u2212 v\u2080",
                "v = v\u2080 + at",
            ]
        },
        {
            "letter": "B",
            "title": "Deriving Kinematic Equation 4",
            "goal": "x = x\u2080 + \u00bd(v\u2080 + v)t",
            "steps": [
                "Average velocity is defined as total displacement over total time. Write the definition as a formula using v\u0304, \u0394x, and \u0394t.",
                "Displacement is \u0394x = x \u2212 x\u2080 and \u0394t = t. Substitute these into the formula you wrote above.",
                "Multiply both sides of your equation by t to clear the denominator. Write the result.",
                "Add x\u2080 to both sides to isolate x on one side. Write the result.",
                "Under constant acceleration, velocity changes linearly. The average of a linearly changing quantity equals the mean of its start and end values. Express v\u0304 as the mean of v\u2080 and v.",
                "In your equation from Step 4, replace v\u0304 with the expression you found in Step 5. Write the final equation.",
            ],
            "answers": [
                "v\u0304 = \u0394x / \u0394t",
                "v\u0304 = (x \u2212 x\u2080) / t",
                "v\u0304 \u00b7 t = x \u2212 x\u2080",
                "x = x\u2080 + v\u0304 \u00b7 t",
                "v\u0304 = (v\u2080 + v) / 2",
                "x = x\u2080 + \u00bd(v\u2080 + v)t",
            ]
        },
        {
            "letter": "C",
            "title": "Deriving Kinematic Equation 2",
            "goal": "x = x\u2080 + v\u2080t + \u00bdat\u00b2",
            "steps": [
                "Write down Equation 1: v = v\u2080 + at.",
                "Write down Equation 4: x = x\u2080 + \u00bd(v\u2080 + v)t.",
                "In Equation 4, replace v with the right-hand side of Equation 1. Write the full substitution without simplifying.",
                "Inside the brackets, combine the like terms v\u2080 + v\u2080 into a single term. Write the simplified expression.",
                "Distribute t into the bracket. Write the result.",
                "Distribute \u00bd to each term. Simplify the coefficients and write the final equation.",
            ],
            "answers": [
                "v = v\u2080 + at",
                "x = x\u2080 + \u00bd(v\u2080 + v)t",
                "x = x\u2080 + \u00bd(v\u2080 + v\u2080 + at)\u00b7t",
                "x = x\u2080 + \u00bd(2v\u2080 + at)\u00b7t",
                "x = x\u2080 + \u00bd(2v\u2080t + at\u00b2)",
                "x = x\u2080 + v\u2080t + \u00bdat\u00b2",
            ]
        },
        {
            "letter": "D",
            "title": "Deriving Kinematic Equation 3",
            "goal": "v\u00b2 = v\u2080\u00b2 + 2a(x \u2212 x\u2080)",
            "steps": [
                "Start with Equation 1: v = v\u2080 + at. Subtract v\u2080 from both sides. Write the result.",
                "Divide both sides of your equation by a to isolate t. Write t = ...",
                "Write the displacement equation from Equation 4: x \u2212 x\u2080 = \u00bd(v\u2080 + v)\u00b7t.",
                "Replace t with the expression you found in Step 2. Write the full substitution.",
                "Multiply the two factors in the numerator using the difference-of-squares identity: (v\u2080 + v)(v \u2212 v\u2080) = v\u00b2 \u2212 v\u2080\u00b2. Simplify.",
                "Multiply both sides by 2a to clear the denominator. Write the result.",
                "Add v\u2080\u00b2 to both sides to isolate v\u00b2. Write the final equation.",
            ],
            "answers": [
                "v \u2212 v\u2080 = at",
                "t = (v \u2212 v\u2080) / a",
                "x \u2212 x\u2080 = \u00bd(v\u2080 + v)\u00b7t",
                "x \u2212 x\u2080 = \u00bd(v\u2080 + v)(v \u2212 v\u2080) / a",
                "x \u2212 x\u2080 = (v\u00b2 \u2212 v\u2080\u00b2) / 2a",
                "2a(x \u2212 x\u2080) = v\u00b2 \u2212 v\u2080\u00b2",
                "v\u00b2 = v\u2080\u00b2 + 2a(x \u2212 x\u2080)",
            ]
        },
    ]


# ---------------------------------------------------------------
# Document assembly
# ---------------------------------------------------------------
def create_worksheet():
    doc = Document()
    parts = get_parts()

    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = PAGE_W
    section.page_height = PAGE_H
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN

    # ============================================
    # PAGE 1: STUDENT WORKSHEET (BLANK) -- table fills nearly the whole page
    # ============================================
    usable_height = PAGE_H - MARGIN - MARGIN  # 7.9in
    build_table(doc, parts, is_answer_key=False, row_height=Inches(7.85), target_pt=WORKSHEET_TARGET_PT)

    # ============================================
    # PAGE 2: ANSWER KEY (identical layout, answers in red)
    # ============================================
    doc.add_page_break()

    ak_label = doc.add_paragraph()
    ak_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(ak_label, before=0, after=2)
    add_run(ak_label, "TEACHER\u2019S ANSWER KEY", bold=True, size=Pt(10), color=TEXT_RED)

    build_table(doc, parts, is_answer_key=True, row_height=Inches(7.55), target_pt=ANSWERKEY_TARGET_PT)

    doc.save('/projects/sandbox/physics/kinematics-worksheet-4.docx')
    print("Worksheet saved: kinematics-worksheet-4.docx")


if __name__ == "__main__":
    create_worksheet()
