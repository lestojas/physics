#!/usr/bin/env python3
"""
Class Participation Tally Sheet generator.

Produces a single .xlsx workbook with two worksheets:
  1. "Participation Tally"  - landscape, 8.5 x 13 in (Folio)
  2. "Class Rep of the Day"  - portrait,  8.5 x 13 in (Folio)

Design goals: modern, clean, professional. Colours use a navy + teal palette
with subtle zebra striping and coloured section bands for MALE / FEMALE.

Run:  python3 tools/generate_tally.py
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# Configuration
# --------------------------------------------------------------------------- #
OUTPUT = "Class_Participation_Tally_G12-Tesla.xlsx"

CLASS_INFO = {
    "title": "CLASS PARTICIPATION TALLY",
    "grade": "Grade 12 - TESLA",
    "subject": "Physics 1",
    "term": "Term 1",
    "teacher": "Engr. Philip Jayson L. Lestojas",
}

MALE = [
    "CALLEDO, VEN GABRIEL D.",
    "CANDOL, CARL HIENZ R.",
    "CUAS, VINCENT LORENCE D.",
    "DEVIBAR, PRINCE ALVIN D.",
    "ERICO, LEHYAN M.",
    "MAHINAY, JULIUS II A.",
    "MORALES, GERALD ANTHONY E.",
    "QUISADA, PRINCE MYKO G.",
    "RODILLA, EZEKIEL V.",
    "TALE, MARK PATRICK D.",
    "SINGSON, CYRIL M.",
]

FEMALE = [
    "ADORMEO, COLLEEN D.",
    "AGUILAR, JENILYNE S.",
    "ALAMANSA, JYKA JADE",
    "ALESNA, JAKAMIA",
    "ANTIPUESTO, CHARLYN A.",
    "ARANCES, DECYRIE I.",
    "BAGINDA, NOR MYNA A.",
    "BUSTAMANTE, JENILYN T.",
    "CALAGO, CHELSEA FLUER LOUIZE M.",
    "CRETA, BENCH TRIXIE B.",
    "CRODUA, CLAIRE E.",
    "CULLAMAT, RUTH ALTHEA T.",
    "DARUNDAY, MELANIE Z.",
    "DOMAGONOT, NEYREN GLACE P.",
    "EDUSMA, NICE B.",
    "GAMARCHA, TONETH FAITH B.",
    "GRAYAN, JONI MARIE E.",
    "HELT, ALYANNA JOY T.",
    "HIPONIA, MARY ROSE S.",
    "LINDO, EDEN MARIE D.",
    "MANLAPAZ, KIRSTEL",
    "OMAS-AS, RHIANNELLE C.",
    "PACQUIAO, JULIA ANGELLI A.",
    "RIVERA, ALLYSA RHEA M.",
    "RODRIGUEZ, YANARA A.",
    "SOLIVA, KYLA JOY R.",
    "SOSMENA, NINA MERIEL O.".replace("SOSMENA", "SOSME\u00d1A").replace("NINA", "NI\u00d1A"),
    "TAMPAN, JHAMAICA M.",
    "VALDAZO, ANGEL P.",
]

NUM_DATES = 12          # number of session/date columns on the tally sheet
NUM_REP_ROWS = 30       # number of rows on the class-representative sheet

# --------------------------------------------------------------------------- #
# Palette
# --------------------------------------------------------------------------- #
# Shared (palette-independent) colours
WHITE      = "FFFFFF"
BORDER_CLR = "AEBECD"   # light grid line
TEXT_DK    = "1F2A37"   # body text

# Selectable colour schemes. Each maps the palette-specific roles.
PALETTES = {
    # Original navy + teal scheme (used by the Physics 1 sheet)
    "navy": {
        "NAVY_DK":    "16304D",   # title band
        "NAVY":       "1F3A5F",   # header band
        "TEAL":       "2E86AB",   # accent band / subtitle
        "LIGHT":      "EAF1F7",   # light fill (date header, teacher line)
        "ZEBRA":      "F4F8FB",   # alternate row
        "MALE_CLR":   "2F5597",   # male section band
        "FEMALE_CLR": "9B3B6E",   # female section band
        "NUMCOL_CLR": "FBEFD6",   # soft amber for the count sub-column
    },
    # Emerald + amber scheme (used by the Research 1 sheet)
    "emerald": {
        "NAVY_DK":    "0F3D2E",   # deep emerald title band
        "NAVY":       "1B5E43",   # emerald header band
        "TEAL":       "C77D0A",   # amber accent band / subtitle
        "LIGHT":      "E6F2EB",   # light green fill
        "ZEBRA":      "F2F9F5",   # alternate row
        "MALE_CLR":   "1B5E43",   # emerald male section band
        "FEMALE_CLR": "B26A00",   # amber-bronze female section band
        "NUMCOL_CLR": "FBEED6",   # soft amber for the count sub-column
    },
}

# Active palette values (populated by apply_palette). Defaults to "navy".
NAVY_DK = NAVY = TEAL = LIGHT = ZEBRA = MALE_CLR = FEMALE_CLR = NUMCOL_CLR = ""


def apply_palette(name):
    """Load a colour scheme into module globals used by the builders."""
    global med
    globals().update(PALETTES[name])
    med = Side(style="medium", color=NAVY)   # rebuild accent border colour

# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #

def in_to_width(inches):
    """Convert a desired print width (inches) to Excel column-width units."""
    return round((inches * 96 - 5) / 7.0, 2)


def fill(color):
    return PatternFill("solid", fgColor=color)


def font(size=10, bold=False, italic=False, color=TEXT_DK, name="Calibri"):
    return Font(name=name, size=size, bold=bold, italic=italic, color=color)


def align(h="center", v="center", wrap=False, rot=0, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap,
                     text_rotation=rot, indent=indent)


thin = Side(style="thin", color=BORDER_CLR)
med = Side(style="medium", color="1F3A5F")   # rebuilt by apply_palette()
ALL_THIN = Border(left=thin, right=thin, top=thin, bottom=thin)


def box(cell, border=ALL_THIN, fillc=None, f=None, a=None, value=None):
    if value is not None:
        cell.value = value
    if border is not None:
        cell.border = border
    if fillc is not None:
        cell.fill = fill(fillc)
    if f is not None:
        cell.font = f
    if a is not None:
        cell.alignment = a
    return cell


# --------------------------------------------------------------------------- #
# Sheet 1 : Participation Tally  (landscape)
# --------------------------------------------------------------------------- #

def build_tally(ws):
    students = [("MALE", None)] + [("M", n) for n in MALE] + \
               [("FEMALE", None)] + [("F", n) for n in FEMALE]

    first_data_date_col = 3                      # column C
    last_col = first_data_date_col + NUM_DATES * 2 - 1
    last_letter = get_column_letter(last_col)

    # ---- column widths -----------------------------------------------------
    ws.column_dimensions["A"].width = in_to_width(0.42)          # No.
    ws.column_dimensions["B"].width = in_to_width(2.30)          # Name
    for d in range(NUM_DATES):
        tally_col = get_column_letter(first_data_date_col + d * 2)
        num_col = get_column_letter(first_data_date_col + d * 2 + 1)
        ws.column_dimensions[tally_col].width = in_to_width(0.52)
        ws.column_dimensions[num_col].width = in_to_width(0.34)

    r = 1
    # ---- Title band --------------------------------------------------------
    ws.merge_cells(f"A{r}:{last_letter}{r}")
    box(ws.cell(r, 1), border=None, fillc=NAVY_DK,
        f=font(20, bold=True, color=WHITE), a=align(),
        value=CLASS_INFO["title"])
    ws.row_dimensions[r].height = 30
    r += 1

    # ---- Subtitle band -----------------------------------------------------
    ws.merge_cells(f"A{r}:{last_letter}{r}")
    subtitle = f"{CLASS_INFO['grade']}     \u2022     {CLASS_INFO['subject']}     \u2022     {CLASS_INFO['term']}"
    box(ws.cell(r, 1), border=None, fillc=TEAL,
        f=font(12, bold=True, color=WHITE), a=align(), value=subtitle)
    ws.row_dimensions[r].height = 20
    r += 1

    # ---- Teacher line ------------------------------------------------------
    ws.merge_cells(f"A{r}:{last_letter}{r}")
    box(ws.cell(r, 1), border=None, fillc=LIGHT,
        f=font(10, italic=True, color=NAVY),
        a=align(h="center"), value=f"Teacher:  {CLASS_INFO['teacher']}")
    ws.row_dimensions[r].height = 16
    r += 1

    # ---- spacer ------------------------------------------------------------
    ws.row_dimensions[r].height = 5
    r += 1

    # ---- Header rows (2) ---------------------------------------------------
    hdr1 = r        # date-entry row (teacher writes date here)
    hdr2 = r + 1    # Tally / # sub-header row

    # No. + Name span both header rows
    ws.merge_cells(f"A{hdr1}:A{hdr2}")
    box(ws.cell(hdr1, 1), fillc=NAVY, f=font(9, bold=True, color=WHITE),
        a=align(wrap=True), value="No.")
    ws.merge_cells(f"B{hdr1}:B{hdr2}")
    box(ws.cell(hdr1, 2), fillc=NAVY, f=font(10, bold=True, color=WHITE),
        a=align(h="center"), value="NAME OF STUDENT")

    for d in range(NUM_DATES):
        c0 = first_data_date_col + d * 2
        c1 = c0 + 1
        # date-entry cell (merged pair, left blank on purpose)
        ws.merge_cells(start_row=hdr1, start_column=c0, end_row=hdr1, end_column=c1)
        box(ws.cell(hdr1, c0), fillc=LIGHT,
            f=font(8, bold=True, color=NAVY), a=align())
        # sub-headers
        box(ws.cell(hdr2, c0), fillc=NAVY,
            f=font(7.5, bold=True, color=WHITE), a=align(), value="Tally")
        box(ws.cell(hdr2, c1), fillc=TEAL,
            f=font(8, bold=True, color=WHITE), a=align(), value="#")

    ws.row_dimensions[hdr1].height = 22
    ws.row_dimensions[hdr2].height = 15
    r = hdr2 + 1

    # ---- Student rows ------------------------------------------------------
    counter = 0
    zebra_toggle = 0
    for kind, name in students:
        if kind in ("MALE", "FEMALE"):
            band = MALE_CLR if kind == "MALE" else FEMALE_CLR
            ws.merge_cells(f"A{r}:{last_letter}{r}")
            box(ws.cell(r, 1), fillc=band,
                f=font(9.5, bold=True, color=WHITE),
                a=align(h="left", indent=1), value=f"  {kind}")
            ws.row_dimensions[r].height = 15
            zebra_toggle = 0
            r += 1
            continue

        counter += 1
        rowfill = ZEBRA if zebra_toggle % 2 else WHITE
        zebra_toggle += 1

        box(ws.cell(r, 1), fillc=rowfill, f=font(9, color=TEXT_DK),
            a=align(), value=counter)
        box(ws.cell(r, 2), fillc=rowfill, f=font(9.5, color=TEXT_DK),
            a=align(h="left", indent=1), value=name)
        for d in range(NUM_DATES):
            c0 = first_data_date_col + d * 2
            c1 = c0 + 1
            box(ws.cell(r, c0), fillc=rowfill)                 # tally (blank)
            box(ws.cell(r, c1), fillc=NUMCOL_CLR)              # count (soft amber)
        ws.row_dimensions[r].height = 14
        r += 1

    table_bottom = r - 1

    # medium outer border around the whole table
    _outline(ws, hdr1, 1, table_bottom, last_col)

    # ---- legend + footer ---------------------------------------------------
    r += 1
    ws.merge_cells(f"A{r}:{last_letter}{r}")
    box(ws.cell(r, 1), border=None,
        f=font(8, italic=True, color="6B7683"), a=align(h="left"),
        value=("Legend:  Write the session date in each shaded header cell.  "
               "Use the wide column for tally marks (e.g. IIII) and the narrow # column for the total count."))
    ws.row_dimensions[r].height = 14

    r += 2
    _signature_block(ws, r, 2, CLASS_INFO["teacher"])

    # ---- print setup -------------------------------------------------------
    ws.print_area = f"A1:{last_letter}{r + 2}"
    ws.sheet_view.showGridLines = False
    ps = ws.page_setup
    ps.orientation = "landscape"
    ps.paperSize = 14                    # 8.5 x 13 in (Folio)
    ps.fitToWidth = 1
    ps.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.3
    ws.page_margins.header = ws.page_margins.footer = 0.1
    ws.freeze_panes = f"C{hdr2 + 1}"


# --------------------------------------------------------------------------- #
# Sheet 2 : Class Representative of the Day  (portrait)
# --------------------------------------------------------------------------- #

def build_reps(ws):
    last_col = 4
    last_letter = get_column_letter(last_col)

    ws.column_dimensions["A"].width = in_to_width(0.5)    # No.
    ws.column_dimensions["B"].width = in_to_width(1.5)    # Date
    ws.column_dimensions["C"].width = in_to_width(3.4)    # Name
    ws.column_dimensions["D"].width = in_to_width(2.6)    # Signature

    r = 1
    # Title
    ws.merge_cells(f"A{r}:{last_letter}{r}")
    box(ws.cell(r, 1), border=None, fillc=NAVY_DK,
        f=font(16, bold=True, color=WHITE), a=align(),
        value="CLASS REPRESENTATIVE OF THE DAY")
    ws.row_dimensions[r].height = 28
    r += 1

    ws.merge_cells(f"A{r}:{last_letter}{r}")
    subtitle = f"{CLASS_INFO['grade']}     \u2022     {CLASS_INFO['subject']}     \u2022     {CLASS_INFO['term']}"
    box(ws.cell(r, 1), border=None, fillc=TEAL,
        f=font(11, bold=True, color=WHITE), a=align(), value=subtitle)
    ws.row_dimensions[r].height = 19
    r += 1

    ws.merge_cells(f"A{r}:{last_letter}{r}")
    box(ws.cell(r, 1), border=None, fillc=LIGHT,
        f=font(10, italic=True, color=NAVY), a=align(),
        value=f"Teacher:  {CLASS_INFO['teacher']}")
    ws.row_dimensions[r].height = 16
    r += 1

    ws.row_dimensions[r].height = 6
    r += 1

    # header
    hdr = r
    headers = ["No.", "DATE", "NAME OF CLASS REPRESENTATIVE", "SIGNATURE"]
    for c, text in enumerate(headers, start=1):
        box(ws.cell(hdr, c), fillc=NAVY, f=font(10, bold=True, color=WHITE),
            a=align(), value=text)
    ws.row_dimensions[hdr].height = 22
    r += 1

    for i in range(NUM_REP_ROWS):
        rowfill = ZEBRA if i % 2 else WHITE
        box(ws.cell(r, 1), fillc=rowfill, f=font(9, color=TEXT_DK),
            a=align(), value=i + 1)
        box(ws.cell(r, 2), fillc=rowfill)                       # date (blank)
        box(ws.cell(r, 3), fillc=rowfill)                       # name (blank)
        box(ws.cell(r, 4), fillc=rowfill)                       # signature (blank)
        ws.row_dimensions[r].height = 20
        r += 1

    table_bottom = r - 1
    _outline(ws, hdr, 1, table_bottom, last_col)

    r += 2
    _signature_block(ws, r, 3, CLASS_INFO["teacher"])

    ws.print_area = f"A1:{last_letter}{r + 2}"
    ws.sheet_view.showGridLines = False
    ps = ws.page_setup
    ps.orientation = "portrait"
    ps.paperSize = 14                    # 8.5 x 13 in (Folio)
    ps.fitToWidth = 1
    ps.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.page_margins.header = ws.page_margins.footer = 0.1


# --------------------------------------------------------------------------- #
# Shared blocks
# --------------------------------------------------------------------------- #

def _outline(ws, r0, c0, r1, c1):
    """Draw a medium navy border around the outer edge of a rectangle."""
    for c in range(c0, c1 + 1):
        top = ws.cell(r0, c)
        bot = ws.cell(r1, c)
        top.border = Border(left=top.border.left, right=top.border.right,
                            top=med, bottom=top.border.bottom)
        bot.border = Border(left=bot.border.left, right=bot.border.right,
                            top=bot.border.top, bottom=med)
    for rr in range(r0, r1 + 1):
        lft = ws.cell(rr, c0)
        rgt = ws.cell(rr, c1)
        lft.border = Border(left=med, right=lft.border.right,
                            top=lft.border.top, bottom=lft.border.bottom)
        rgt.border = Border(left=rgt.border.left, right=med,
                            top=rgt.border.top, bottom=rgt.border.bottom)


def _signature_block(ws, r, col, teacher):
    """Signature line with teacher name and 'Subject Teacher' label."""
    name_cell = ws.cell(r, col)
    name_cell.value = teacher
    name_cell.font = font(11, bold=True, color=TEXT_DK)
    name_cell.alignment = align(h="center")
    name_cell.border = Border(top=Side(style="thin", color=TEXT_DK))
    ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col + 1)
    ws.cell(r, col + 1).border = Border(top=Side(style="thin", color=TEXT_DK))

    lbl = ws.cell(r + 1, col)
    lbl.value = "Subject Teacher"
    lbl.font = font(10, italic=True, color="6B7683")
    lbl.alignment = align(h="center")
    ws.merge_cells(start_row=r + 1, start_column=col, end_row=r + 1, end_column=col + 1)


# --------------------------------------------------------------------------- #
def build_workbook(output, subject, palette):
    apply_palette(palette)
    CLASS_INFO["subject"] = subject

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Participation Tally"
    build_tally(ws1)

    ws2 = wb.create_sheet("Class Rep of the Day")
    build_reps(ws2)

    wb.save(output)
    print(f"Saved {output}  (subject='{subject}', palette='{palette}')")


def main():
    import argparse
    p = argparse.ArgumentParser(description="Generate a class participation tally workbook.")
    p.add_argument("--output", default=OUTPUT, help="output .xlsx filename")
    p.add_argument("--subject", default=CLASS_INFO["subject"], help="subject label in the header")
    p.add_argument("--palette", default="navy", choices=list(PALETTES),
                   help="colour scheme")
    args = p.parse_args()
    build_workbook(args.output, args.subject, args.palette)


if __name__ == "__main__":
    main()
