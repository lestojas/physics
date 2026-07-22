#!/usr/bin/env python3
"""
Faithful PNG preview of an .xlsx file.

Reads the *actual* generated workbook back with openpyxl (values, merges,
fills, fonts, alignment, borders, column widths, row heights) and renders
each worksheet to a PNG using Pillow. This lets you review the real layout
before finalising / sharing the Excel file.

Run:  python3 tools/preview_xlsx.py Class_Participation_Tally_G12-Tesla.xlsx
"""

import sys
import os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from PIL import Image, ImageDraw, ImageFont

FONT_DIR = "/usr/share/fonts/google-noto"
FONTS = {
    (False, False): f"{FONT_DIR}/NotoSans-Regular.ttf",
    (True, False):  f"{FONT_DIR}/NotoSans-Bold.ttf",
    (False, True):  f"{FONT_DIR}/NotoSans-Italic.ttf",
    (True, True):   f"{FONT_DIR}/NotoSans-BoldItalic.ttf",
}
SCALE = 2                     # supersampling for crisp text
DEFAULT_COL_W = 8.43
DEFAULT_ROW_H = 15.0

_font_cache = {}


def get_font(size_pt, bold, italic):
    key = (round(size_pt * SCALE), bold, italic)
    if key not in _font_cache:
        px = max(6, round(size_pt * SCALE * 96 / 72))
        _font_cache[key] = ImageFont.truetype(FONTS[(bold, italic)], px)
    return _font_cache[key]


def col_px(width):
    return round((width if width else DEFAULT_COL_W) * 7) + 5


def row_px(height):
    return round((height if height else DEFAULT_ROW_H) * 96 / 72)


def argb_to_rgb(color, default=None):
    """Return an (r,g,b) tuple from an openpyxl color, or `default`."""
    if color is None:
        return default
    rgb = getattr(color, "rgb", None)
    if not isinstance(rgb, str):
        return default
    if len(rgb) == 8:
        rgb = rgb[2:]
    if len(rgb) != 6:
        return default
    try:
        return tuple(int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return default


def render_sheet(ws, out_path):
    # -- used range -----------------------------------------------------------
    max_row, max_col = ws.max_row, ws.max_column

    # column x-edges
    xs = [0]
    for c in range(1, max_col + 1):
        dim = ws.column_dimensions.get(get_column_letter(c))
        w = dim.width if dim and dim.width else DEFAULT_COL_W
        xs.append(xs[-1] + col_px(w) * SCALE)
    # row y-edges
    ys = [0]
    for r in range(1, max_row + 1):
        dim = ws.row_dimensions.get(r)
        h = dim.height if dim and dim.height else DEFAULT_ROW_H
        ys.append(ys[-1] + row_px(h) * SCALE)

    W, H = xs[-1], ys[-1]
    img = Image.new("RGB", (W + 1, H + 1), "white")
    draw = ImageDraw.Draw(img)

    # -- merged cell handling -------------------------------------------------
    covered = {}          # (r,c) -> anchor (r,c)
    span = {}             # anchor (r,c) -> (r1,c1)
    for mr in ws.merged_cells.ranges:
        c0, r0, c1, r1 = range_boundaries(str(mr))
        span[(r0, c0)] = (r1, c1)
        for rr in range(r0, r1 + 1):
            for cc in range(c0, c1 + 1):
                if (rr, cc) != (r0, c0):
                    covered[(rr, cc)] = (r0, c0)

    def cell_rect(r, c):
        r1, c1 = span.get((r, c), (r, c))
        return xs[c - 1], ys[r - 1], xs[c1], ys[r1]

    # -- pass 1: fills --------------------------------------------------------
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(r, c)
            if cell.fill and cell.fill.patternType == "solid":
                rgb = argb_to_rgb(cell.fill.fgColor)
                if rgb:
                    x0, y0, x1, y1 = cell_rect(r, c)
                    draw.rectangle([x0, y0, x1, y1], fill=rgb)

    # -- pass 2: borders ------------------------------------------------------
    def line(p0, p1, side):
        if side is None or side.style is None:
            return
        color = argb_to_rgb(side.color, default=(120, 130, 140))
        w = 2 * SCALE if side.style in ("medium", "thick") else 1 * SCALE
        draw.line([p0, p1], fill=color, width=w)

    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(r, c)
            b = cell.border
            if not b:
                continue
            x0, y0, x1, y1 = cell_rect(r, c)
            line((x0, y0), (x1, y0), b.top)
            line((x0, y1), (x1, y1), b.bottom)
            line((x0, y0), (x0, y1), b.left)
            line((x1, y0), (x1, y1), b.right)

    # -- pass 3: text ---------------------------------------------------------
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(r, c)
            if cell.value is None or cell.value == "":
                continue
            text = str(cell.value)
            f = cell.font
            size = f.size or 11
            bold = bool(f.bold)
            italic = bool(f.italic)
            fnt = get_font(size, bold, italic)
            color = argb_to_rgb(f.color, default=(20, 30, 40))

            x0, y0, x1, y1 = cell_rect(r, c)
            al = cell.alignment
            halign = al.horizontal or "left"
            valign = al.vertical or "bottom"
            indent = (al.indent or 0) * 7 * SCALE

            bbox = draw.textbbox((0, 0), text, font=fnt)
            tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
            pad = 3 * SCALE

            if halign == "center":
                tx = x0 + (x1 - x0 - tw) / 2
            elif halign == "right":
                tx = x1 - tw - pad
            else:
                tx = x0 + pad + indent
            if valign == "center":
                ty = y0 + (y1 - y0 - th) / 2 - bbox[1]
            elif valign == "top":
                ty = y0 + pad - bbox[1]
            else:
                ty = y1 - th - pad - bbox[1]

            draw.text((tx, ty), text, font=fnt, fill=color)

    # -- downscale for smooth output -----------------------------------------
    final = img.resize((max(1, (W + 1) // SCALE), max(1, (H + 1) // SCALE)),
                       Image.LANCZOS)
    final.save(out_path)
    print(f"  wrote {out_path}  ({final.width}x{final.height}px)")


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else "Class_Participation_Tally_G12-Tesla.xlsx"
    wb = load_workbook(path)
    os.makedirs("preview", exist_ok=True)
    stem = os.path.splitext(os.path.basename(path))[0]
    for i, ws in enumerate(wb.worksheets, start=1):
        safe = ws.title.replace(" ", "_")
        render_sheet(ws, f"preview/{stem}__{i}_{safe}.png")


if __name__ == "__main__":
    main()
