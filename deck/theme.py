"""
Design system + layout helpers for the Lecture 3 PowerPoint deck.

Modern, creative 16:9 theme built around the coastal-ocean case study that
runs through the whole lecture. All slides are constructed on blank layouts so
we have full control over typography, colour and diagram geometry.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
import copy

# ----------------------------------------------------------------------------
# Canvas (16:9 widescreen)
# ----------------------------------------------------------------------------
EMU_W = Inches(13.333)
EMU_H = Inches(7.5)

# ----------------------------------------------------------------------------
# Palette
# ----------------------------------------------------------------------------
INK        = RGBColor(0x14, 0x25, 0x2E)   # near-black ink
INK_SOFT   = RGBColor(0x3A, 0x4B, 0x54)   # muted body text
PAPER      = RGBColor(0xF6, 0xF9, 0xFA)   # off-white background
PAPER_2    = RGBColor(0xEC, 0xF2, 0xF3)   # panel background
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
CLOUD      = RGBColor(0xDD, 0xE7, 0xE9)   # hairline / border

DEEP       = RGBColor(0x0C, 0x3B, 0x47)   # deep ocean (primary dark)
TEAL       = RGBColor(0x14, 0x82, 0x8C)   # teal
AQUA       = RGBColor(0x2A, 0x9D, 0x8F)   # aqua green
SKY        = RGBColor(0x2B, 0x6C, 0xB2)   # blue
SAND       = RGBColor(0xE9, 0xC4, 0x6A)   # sand / gold
CORAL      = RGBColor(0xE7, 0x6F, 0x51)   # terracotta / coral
PLUM       = RGBColor(0x7B, 0x4E, 0x8C)   # purple
LEAF       = RGBColor(0x43, 0x8A, 0x5E)   # green
RED        = RGBColor(0xC6, 0x4A, 0x3F)   # highlight red

# Section accents (index by section key)
SECTIONS = {
    "OPENER": ("The Research Problem Journey", DEEP),
    "P1":     ("Part 1 - Evaluating Sources", TEAL),
    "P2":     ("Part 2 - Synthesizing Literature", SKY),
    "P3":     ("Part 3 - The Problem Statement", CORAL),
    "P4":     ("Part 4 - Questions & Hypotheses", PLUM),
    "P5":     ("Part 5 - Justifying with CER", LEAF),
    "P6":     ("Part 6 - Assumptions & Limitations", SAND),
    "CLOSE":  ("The Full Journey", DEEP),
}

# Source-key pill styling
SOURCE_STYLES = {
    "textbook":   ("TEXTBOOK-BASED", SKY, WHITE),
    "guide":      ("GUIDE-BASED", CORAL, WHITE),
    "instructor": ("INSTRUCTOR-CREATED", AQUA, WHITE),
    "mixed":      ("TEXTBOOK + GUIDE", PLUM, WHITE),
    "tb_guide":   ("TEXTBOOK + GUIDE", PLUM, WHITE),
    "tb_inst":    ("TEXTBOOK + INSTRUCTOR", PLUM, WHITE),
}

FONT      = "Segoe UI"
FONT_LIGHT = "Segoe UI Light"
FONT_SEMI = "Segoe UI Semibold"


# ----------------------------------------------------------------------------
# Low-level helpers
# ----------------------------------------------------------------------------
def _set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _no_line(shape):
    shape.line.fill.background()


def _line(shape, color, w=1.0):
    shape.line.color.rgb = color
    shape.line.width = Pt(w)


def _no_shadow(shape):
    # remove the default preset shadow that python-pptx auto shapes carry
    sp = shape._element.spPr
    existing = sp.find(qn('a:effectLst'))
    if existing is None:
        sp.append(sp.makeelement(qn('a:effectLst'), {}))


def _soft_shadow(shape, blur=0.07, dist=0.045, alpha=72):
    """Add a soft drop shadow to a shape."""
    sp = shape._element.spPr
    for tag in ('a:effectLst',):
        e = sp.find(qn(tag))
        if e is not None:
            sp.remove(e)
    eff = sp.makeelement(qn('a:effectLst'), {})
    sh = eff.makeelement(qn('a:outerShdw'), {
        'blurRad': str(Emu(Inches(blur)).emu if hasattr(Emu(Inches(blur)), 'emu') else int(Inches(blur))),
        'dist': str(int(Inches(dist))),
        'dir': '5400000', 'rotWithShape': '0'})
    clr = sh.makeelement(qn('a:srgbClr'), {'val': '1E2A30'})
    a = clr.makeelement(qn('a:alpha'), {'val': str(int((100 - alpha) * 1000))})
    clr.append(a)
    sh.append(clr)
    eff.append(sh)
    sp.append(eff)


def rect(slide, x, y, w, h, color=None, shape=MSO_SHAPE.RECTANGLE,
         line_color=None, line_w=1.0, shadow=False):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if color is not None:
        _set_fill(sp, color)
    else:
        sp.fill.background()
    if line_color is not None:
        _line(sp, line_color, line_w)
    else:
        _no_line(sp)
    if shadow:
        _soft_shadow(sp)
    else:
        _no_shadow(sp)
    sp.shadow.inherit = False
    return sp


def _apply_runs(tf, para, text, size, color, bold, italic, font=FONT):
    """Support simple **bold** and *italic* inline markup within a paragraph."""
    import re
    tokens = re.split(r'(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*)', text)
    first = True
    for tok in tokens:
        if tok == '':
            continue
        b, it = bold, italic
        t = tok
        if tok.startswith('***') and tok.endswith('***'):
            b = True
            it = True
            t = tok[3:-3]
        elif tok.startswith('**') and tok.endswith('**'):
            b = True
            t = tok[2:-2]
        elif tok.startswith('*') and tok.endswith('*'):
            it = True
            t = tok[1:-1]
        run = para.add_run()
        run.text = t
        run.font.size = Pt(size)
        run.font.bold = b
        run.font.italic = it
        run.font.name = font
        run.font.color.rgb = color
    return para


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, wrap=True):
    """lines: list of dicts {text,size,color,bold,italic,space_after,
    space_before,bullet,level,line_spacing,font}."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", align)
        if ln.get("space_after") is not None:
            p.space_after = Pt(ln["space_after"])
        if ln.get("space_before") is not None:
            p.space_before = Pt(ln["space_before"])
        if ln.get("line_spacing") is not None:
            p.line_spacing = ln["line_spacing"]
        _apply_runs(tf, p, ln["text"], ln.get("size", 18),
                    ln.get("color", INK), ln.get("bold", False),
                    ln.get("italic", False), ln.get("font", FONT))
    return tb


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text



# ----------------------------------------------------------------------------
# Slide scaffolding
# ----------------------------------------------------------------------------
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = EMU_W
        self.prs.slide_height = EMU_H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def save(self, path):
        self.prs.save(path)

    def new(self):
        return self.prs.slides.add_slide(self.blank)


def base(slide, section, decor_waves=True):
    """Paint the standard content-slide background + section spine."""
    color = SECTIONS[section][1]
    rect(slide, 0, 0, 13.333, 7.5, PAPER)                     # page
    rect(slide, 0, 0, 0.28, 7.5, color)                       # left spine
    # faint corner accent circle bottom-right
    if decor_waves:
        c = rect(slide, 11.9, 6.35, 2.6, 2.6, None,
                 shape=MSO_SHAPE.OVAL)
        c.fill.solid(); c.fill.fore_color.rgb = color
        c.line.fill.background()
        _set_transparency(c, 90)
    return color


def _set_transparency(shape, pct):
    """Set solid-fill transparency (pct = how transparent)."""
    sp = shape.fill.fore_color._xFill
    srgb = sp.find(qn('a:srgbClr'))
    if srgb is None:
        return
    a = srgb.makeelement(qn('a:alpha'), {'val': str(int((100 - pct) * 1000))})
    srgb.append(a)


def header(slide, section, title, subtitle=None, source=None, number=None):
    color = SECTIONS[section][1]
    tag = SECTIONS[section][0]
    # section kicker
    textbox(slide, 0.62, 0.42, 9.5, 0.35,
            [{"text": tag.upper(), "size": 12.5, "color": color, "bold": True}])
    # title
    textbox(slide, 0.6, 0.74, 10.7, 1.15,
            [{"text": title, "size": 30, "color": INK, "bold": True,
              "line_spacing": 0.98}])
    y = 1.86
    if subtitle:
        textbox(slide, 0.62, y, 11.5, 0.5,
                [{"text": subtitle, "size": 15, "color": INK_SOFT,
                  "italic": True}])
        y += 0.5
    # accent rule under header
    rect(slide, 0.63, 1.72, 1.7, 0.055, color)
    if source:
        source_pill(slide, source)
    if number is not None:
        footer(slide, section, number)
    return color


def footer(slide, section, number):
    color = SECTIONS[section][1]
    textbox(slide, 11.6, 7.02, 1.4, 0.35,
            [{"text": f"{number:02d}  /  51", "size": 10.5,
              "color": INK_SOFT, "bold": True, "align": PP_ALIGN.RIGHT}])
    textbox(slide, 0.62, 7.02, 8.0, 0.35,
            [{"text": "Lecture 3  -  Review of Literature for Identifying Research Problems",
              "size": 9.5, "color": CLOUD if False else INK_SOFT}])


def source_pill(slide, source):
    label, fill, fg = SOURCE_STYLES[source]
    w = 0.115 * len(label) + 0.45
    x = 13.333 - 0.6 - w
    pill = rect(slide, x, 0.46, w, 0.34, fill,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _radius(pill, 0.5)
    textbox(slide, x, 0.46, w, 0.34,
            [{"text": label, "size": 10.5, "color": fg, "bold": True,
              "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)


def _radius(shape, frac):
    """Set corner radius fraction on a rounded rectangle."""
    try:
        shape.adjustments[0] = frac
    except Exception:
        pass


# ----------------------------------------------------------------------------
# Content components
# ----------------------------------------------------------------------------
def bullets(slide, x, y, w, h, items, color, size=18, gap=9):
    """items: list of dicts {text, level(0/1), bold, italic} using inline
    markup. Level-0 bullets get a coloured square marker."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        lvl = it.get("level", 0)
        p.space_after = Pt(it.get("gap", gap))
        p.line_spacing = it.get("line_spacing", 1.04)
        indent = 0.0 if lvl == 0 else 0.42
        # marker
        mk = p.add_run()
        if lvl == 0:
            mk.text = "\u25A0  "
            mk.font.color.rgb = color
            mk.font.size = Pt(size - 4)
        else:
            mk.text = "\u2013  "
            mk.font.color.rgb = INK_SOFT
            mk.font.size = Pt(size - 2)
        p.level = lvl
        _apply_runs(tf, p, it["text"], it.get("size", size - lvl * 1),
                    it.get("color", INK if lvl == 0 else INK_SOFT),
                    it.get("bold", False), it.get("italic", False))
    return tb


def numbered(slide, x, y, w, h, items, color, size=18, gap=10):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.line_spacing = 1.04
        num = p.add_run()
        num.text = f"{it.get('n', i+1)}.  "
        num.font.color.rgb = color
        num.font.bold = True
        num.font.size = Pt(size)
        _apply_runs(tf, p, it["text"], size, it.get("color", INK),
                    it.get("bold", False), it.get("italic", False))
    return tb


def card(slide, x, y, w, h, fill=WHITE, line=CLOUD, radius=0.08, shadow=True):
    c = rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             line_color=line, line_w=1.0, shadow=shadow)
    _radius(c, radius)
    return c


def table(slide, x, y, w, headers, rows, accent, col_widths=None,
          header_size=13.5, body_size=12, row_h=0.5, header_h=0.5,
          first_col_bold=False):
    """Draw a styled table. headers/rows contain inline-markup strings."""
    ncol = len(headers)
    nrow = len(rows) + 1
    total_h = header_h + row_h * len(rows)
    gt = slide.shapes.add_table(nrow, ncol, Inches(x), Inches(y),
                                Inches(w), Inches(header_h + row_h * len(rows)))
    tbl = gt.table
    # disable banded styling; we colour manually
    tbl.first_row = False
    tbl.horz_banding = False
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(cw)
    tbl.rows[0].height = Inches(header_h)
    for r in range(1, nrow):
        tbl.rows[r].height = Inches(row_h)

    def style_cell(cell, text, size, color, bold, fill, align=PP_ALIGN.LEFT,
                   italic=False):
        cell.fill.solid(); cell.fill.fore_color.rgb = fill
        cell.margin_left = Inches(0.11); cell.margin_right = Inches(0.09)
        cell.margin_top = Inches(0.05); cell.margin_bottom = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        # clear default
        p.text = ""
        _apply_runs(tf, p, text if text else "", size, color, bold, italic)

    for c in range(ncol):
        style_cell(tbl.cell(0, c), headers[c], header_size, WHITE, True,
                   accent, PP_ALIGN.LEFT)
    for r, row in enumerate(rows, start=1):
        rowfill = WHITE if r % 2 == 1 else PAPER_2
        for c in range(ncol):
            bold = first_col_bold and c == 0
            style_cell(tbl.cell(r, c), row[c], body_size,
                       INK if c == 0 else INK_SOFT, bold, rowfill)
    return gt



# ----------------------------------------------------------------------------
# Diagram components
# ----------------------------------------------------------------------------
def arrow_between(slide, x, y, w, h, color, direction="right"):
    shp = {
        "right": MSO_SHAPE.RIGHT_ARROW,
        "down": MSO_SHAPE.DOWN_ARROW,
    }[direction]
    a = rect(slide, x, y, w, h, color, shape=shp)
    return a


def chevron_flow(slide, x, y, w, h, nodes, color, num_color=WHITE,
                 label_size=13):
    """Horizontal sequence of numbered chevron nodes."""
    n = len(nodes)
    gap = 0.12
    cw = (w - gap * (n - 1)) / n
    alt = _mix(color, WHITE, 0.26)
    for i, (num, txt) in enumerate(nodes):
        cx = x + i * (cw + gap)
        node_col = color if i % 2 == 0 else alt
        node = rect(slide, cx, y, cw, h, node_col,
                    shape=MSO_SHAPE.CHEVRON, shadow=True)
        # number circle
        circ = rect(slide, cx + 0.12, y + h/2 - 0.19, 0.38, 0.38, WHITE,
                    shape=MSO_SHAPE.OVAL)
        textbox(slide, cx + 0.12, y + h/2 - 0.19, 0.38, 0.38,
                [{"text": str(num), "size": 15, "color": node_col,
                  "bold": True, "align": PP_ALIGN.CENTER}],
                anchor=MSO_ANCHOR.MIDDLE)
        textbox(slide, cx + 0.52, y, cw - 0.62, h,
                [{"text": txt, "size": label_size, "color": WHITE, "bold": True,
                  "line_spacing": 0.95}], anchor=MSO_ANCHOR.MIDDLE)


def step_cards(slide, x, y, w, h, items, color, cols=None, gap=0.22,
               num_size=17, title_size=14, body_size=11.5):
    """Grid of numbered step cards. items = list of (num, title, body)."""
    n = len(items)
    cols = cols or n
    rows = (n + cols - 1) // cols
    cw = (w - gap * (cols - 1)) / cols
    ch = (h - gap * (rows - 1)) / rows
    for i, (num, title, body) in enumerate(items):
        r = i // cols
        c = i % cols
        cx = x + c * (cw + gap)
        cy = y + r * (ch + gap)
        card(slide, cx, cy, cw, ch, WHITE, CLOUD)
        rect(slide, cx, cy, 0.09, ch, color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        badge = rect(slide, cx + 0.22, cy + 0.2, 0.44, 0.44, color,
                     shape=MSO_SHAPE.OVAL)
        textbox(slide, cx + 0.22, cy + 0.2, 0.44, 0.44,
                [{"text": str(num), "size": num_size, "color": WHITE,
                  "bold": True, "align": PP_ALIGN.CENTER}],
                anchor=MSO_ANCHOR.MIDDLE)
        lines = [{"text": title, "size": title_size, "color": INK,
                  "bold": True, "space_after": 4, "line_spacing": 0.98}]
        if body:
            lines.append({"text": body, "size": body_size, "color": INK_SOFT,
                          "line_spacing": 1.0})
        textbox(slide, cx + 0.78, cy + 0.16, cw - 0.95, ch - 0.3, lines,
                anchor=MSO_ANCHOR.MIDDLE)


def chain_vertical(slide, x, y, w, items, color, root_label=None,
                   step_h=0.62, gap=0.22, q_w=None, body_size=12.5):
    """Vertical chain of Q -> A steps with connecting arrows.
    items = list of (question, answer)."""
    q_w = q_w or 1.5
    cy = y
    for i, (q, a) in enumerate(items):
        card(slide, x, cy, w, step_h, WHITE, CLOUD)
        qb = rect(slide, x, cy, q_w, step_h, color,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        textbox(slide, x, cy, q_w, step_h,
                [{"text": q, "size": 12.5, "color": WHITE, "bold": True,
                  "align": PP_ALIGN.CENTER, "line_spacing": 0.95}],
                anchor=MSO_ANCHOR.MIDDLE)
        textbox(slide, x + q_w + 0.18, cy, w - q_w - 0.34, step_h,
                [{"text": a, "size": body_size, "color": INK,
                  "line_spacing": 0.97}], anchor=MSO_ANCHOR.MIDDLE)
        cy += step_h
        if i < len(items) - 1 or root_label:
            arrow_between(slide, x + w/2 - 0.11, cy + 0.01, 0.22, gap - 0.02,
                          color, "down")
            cy += gap
    if root_label:
        rc = rect(slide, x, cy, w, step_h + 0.06, color,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
        textbox(slide, x + 0.2, cy, w - 0.4, step_h + 0.06,
                [{"text": root_label, "size": 13, "color": WHITE, "bold": True,
                  "align": PP_ALIGN.CENTER, "line_spacing": 0.95}],
                anchor=MSO_ANCHOR.MIDDLE)


def funnel(slide, x, y, top_w, bottom_w, total_h, items, color):
    """Stacked narrowing bars (funnel). items = list of (label, body)."""
    n = len(items)
    bar_h = (total_h - 0.16 * (n - 1)) / n
    cy = y
    for i, (label, body) in enumerate(items):
        frac = i / max(n - 1, 1)
        bw = top_w - (top_w - bottom_w) * frac
        cx = x + (top_w - bw) / 2
        shade = color if i == 0 else _mix(color, WHITE, 0.13 * i)
        bar = rect(slide, cx, cy, bw, bar_h, shade,
                   shape=MSO_SHAPE.ROUNDED_RECTANGLE, shadow=True)
        lines = [{"text": label, "size": 15, "color": WHITE, "bold": True,
                  "align": PP_ALIGN.CENTER, "space_after": 2,
                  "line_spacing": 0.95}]
        if body:
            lines.append({"text": body, "size": 11.5, "color": WHITE,
                          "align": PP_ALIGN.CENTER, "line_spacing": 0.95})
        textbox(slide, cx + 0.2, cy, bw - 0.4, bar_h, lines,
                anchor=MSO_ANCHOR.MIDDLE)
        cy += bar_h + 0.16


def _mix(c1, c2, t):
    return RGBColor(
        int(c1[0] + (c2[0] - c1[0]) * t),
        int(c1[1] + (c2[1] - c1[1]) * t),
        int(c1[2] + (c2[2] - c1[2]) * t),
    )


def quad_grid(slide, x, y, w, h, cells, colors, gap=0.2):
    """2x2 grid. cells = list of 4 dicts {title, body}. colors = list of 4."""
    cw = (w - gap) / 2
    ch = (h - gap) / 2
    pos = [(0, 0), (1, 0), (0, 1), (1, 1)]
    for i, (cx_i, cy_i) in enumerate(pos):
        cx = x + cx_i * (cw + gap)
        cy = y + cy_i * (ch + gap)
        card(slide, cx, cy, cw, ch, WHITE, CLOUD)
        rect(slide, cx, cy, cw, 0.12, colors[i],
             shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        lines = [{"text": cells[i]["title"], "size": 15.5, "color": colors[i],
                  "bold": True, "space_after": 5, "line_spacing": 0.98}]
        lines.append({"text": cells[i]["body"], "size": 12.5, "color": INK_SOFT,
                      "line_spacing": 1.02})
        textbox(slide, cx + 0.24, cy + 0.28, cw - 0.44, ch - 0.44, lines,
                anchor=MSO_ANCHOR.MIDDLE)


def icon_checklist(slide, x, y, w, items, color, item_h=0.86, gap=0.14,
                   glyph_size=17):
    """Vertical list of check items with glyph badges.
    items = list of (glyph, title, body)."""
    cy = y
    for glyph, title, body in items:
        card(slide, x, cy, w, item_h, WHITE, CLOUD)
        badge = rect(slide, x + 0.16, cy + item_h/2 - 0.27, 0.54, 0.54, color,
                     shape=MSO_SHAPE.OVAL)
        textbox(slide, x + 0.16, cy + item_h/2 - 0.29, 0.54, 0.54,
                [{"text": glyph, "size": glyph_size, "color": WHITE,
                  "bold": True, "align": PP_ALIGN.CENTER}],
                anchor=MSO_ANCHOR.MIDDLE)
        lines = [{"text": title, "size": 15, "color": INK, "bold": True,
                  "space_after": 2, "line_spacing": 0.95}]
        if body:
            lines.append({"text": body, "size": 12, "color": INK_SOFT,
                          "line_spacing": 0.98})
        textbox(slide, x + 0.86, cy + 0.1, w - 1.05, item_h - 0.2, lines,
                anchor=MSO_ANCHOR.MIDDLE)
        cy += item_h + gap


def cer_boxes(slide, x, y, w, items, colors, box_h=1.3, gap=0.3, horizontal=True):
    """Claim -> Evidence -> Reasoning style linked boxes."""
    n = len(items)
    if horizontal:
        bw = (w - gap * (n - 1)) / n
        for i, (title, body) in enumerate(items):
            cx = x + i * (bw + gap)
            card(slide, cx, y, bw, box_h, WHITE, CLOUD)
            rect(slide, cx, y, bw, 0.42, colors[i],
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            textbox(slide, cx, y, bw, 0.42,
                    [{"text": title, "size": 14, "color": WHITE, "bold": True,
                      "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
            textbox(slide, cx + 0.16, y + 0.5, bw - 0.32, box_h - 0.6,
                    [{"text": body, "size": 11.5, "color": INK_SOFT,
                      "align": PP_ALIGN.CENTER, "line_spacing": 1.0}],
                    anchor=MSO_ANCHOR.MIDDLE)
            if i < n - 1:
                arrow_between(slide, cx + bw + 0.03, y + box_h/2 - 0.14,
                              gap - 0.06, 0.28, colors[i])


def callout(slide, x, y, w, h, text, color, italic=True, size=14):
    """A tinted callout strip for the 'each stop builds' style notes."""
    bg = _mix(color, WHITE, 0.82)
    c = rect(slide, x, y, w, h, bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _radius(c, 0.12)
    rect(slide, x, y, 0.09, h, color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(slide, x + 0.28, y, w - 0.5, h,
            [{"text": text, "size": size, "color": INK, "italic": italic,
              "line_spacing": 1.02}], anchor=MSO_ANCHOR.MIDDLE)


def big_paragraph(slide, x, y, w, h, text, color, size=15.5,
                  quote=True, label_map=None):
    """A framed 'written-out example' paragraph card."""
    card(slide, x, y, w, h, WHITE, CLOUD)
    rect(slide, x, y, 0.11, h, color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    textbox(slide, x + 0.4, y + 0.05, 0.6, 0.7,
            [{"text": "\u201C", "size": 44, "color": _mix(color, WHITE, 0.35),
              "bold": True}])
    textbox(slide, x + 0.42, y + 0.62, w - 0.75, h - 0.85,
            [{"text": text, "size": size, "color": INK, "italic": True,
              "line_spacing": 1.14}], anchor=MSO_ANCHOR.TOP)
