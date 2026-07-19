"""
Minimal OMML (Office Math Markup Language) builder helpers for python-docx.

python-docx has no native support for <m:oMath> equations, so this module
builds the raw OOXML math elements directly via lxml, in the
http://schemas.openxmlformats.org/officeDocument/2006/math namespace,
and provides functions to insert them into paragraphs / table cells.

Reference: ECMA-376 Part 1, section 22.1 (Office Open XML Math / OMML).
"""
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _m(tag):
    """Return the 'm:tag' prefixed name for use with OxmlElement()."""
    return "m:" + tag


def _w(tag):
    """Return the 'w:tag' prefixed name for use with OxmlElement()."""
    return "w:" + tag


def _el(tag, *children, **attrs):
    """Create an OOXML element and append (deep-copied) children.

    Deep-copying every child here -- rather than only in cat() -- is what
    makes it safe to build a fragment once (e.g. `v2 = SUBX(V('v'), '2')`)
    and reuse that same Python list/element across many different
    equations, derivation rows, and table cells throughout the document.
    lxml elements can only live in one place in the XML tree at a time, so
    without this every second use of a reused fragment would silently
    *move* (and corrupt) the first equation that used it.
    """
    import copy as _copy
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(k), v)
    for c in children:
        if c is not None:
            el.append(_copy.deepcopy(c))
    return el


# ---------------------------------------------------------------------------
# Core math runs
# ---------------------------------------------------------------------------

def mrun(text, bold=False, italic=None, color=None):
    """A single <m:r> math run containing literal text.

    Numbers/operators are marked with <m:nor> (normal text, no auto math
    italics) when italic is explicitly False, matching Word's own behavior
    of keeping digits/operators upright while variables stay italic.
    """
    r = _el(_m("r"))
    need_wrpr = bold or color is not None
    need_mrpr = italic is False
    if need_mrpr:
        mrpr = _el(_m("rPr"))
        mrpr.append(_el(_m("nor")))
        r.append(mrpr)
    if need_wrpr:
        wrpr = _el(_w("rPr"))
        if bold:
            wrpr.append(_el(_w("b")))
        if color is not None:
            wrpr.append(_el(_w("color"), **{"w:val": color}))
        r.append(wrpr)
    t = _el(_m("t"))
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    return r


def mtext(text, **kw):
    """Alias for a plain math run."""
    return mrun(text, **kw)


def group(*items):
    """Just returns the list of items (helper for readability)."""
    return [i for i in items if i is not None]


# ---------------------------------------------------------------------------
# Structures: fraction, superscript, subscript, radical, delimiter, box
# ---------------------------------------------------------------------------

def frac(num_items, den_items):
    f = _el(_m("f"))
    num = _el(_m("num"), *num_items)
    den = _el(_m("den"), *den_items)
    f.append(num)
    f.append(den)
    return f


def sup(base_items, sup_items):
    s = _el(_m("sSup"))
    e = _el(_m("e"), *base_items)
    sp = _el(_m("sup"), *sup_items)
    s.append(e)
    s.append(sp)
    return s


def sub(base_items, sub_items):
    s = _el(_m("sSub"))
    e = _el(_m("e"), *base_items)
    sb = _el(_m("sub"), *sub_items)
    s.append(e)
    s.append(sb)
    return s


def sqrt(items, degree_items=None):
    rad = _el(_m("rad"))
    radpr = _el(_m("radPr"))
    if degree_items is None:
        deghide = _el(_m("degHide"), **{"m:val": "1"})
        radpr.append(deghide)
    rad.append(radpr)
    deg = _el(_m("deg"), *(degree_items or []))
    e = _el(_m("e"), *items)
    rad.append(deg)
    rad.append(e)
    return rad


def paren(items, open_ch="(", close_ch=")"):
    d = _el(_m("d"))
    dpr = _el(_m("dPr"))
    dpr.append(_el(_m("begChr"), **{"m:val": open_ch}))
    dpr.append(_el(_m("endChr"), **{"m:val": close_ch}))
    d.append(dpr)
    e = _el(_m("e"), *items)
    d.append(e)
    return d


def boxed(items, color="FF0000"):
    """Rectangle border around an expression (used for final boxed answers)."""
    box = _el(_m("borderBox"))
    e = _el(_m("e"), *items)
    box.append(e)
    return box


# ---------------------------------------------------------------------------
# Wrapping into oMath / oMathPara and inserting into the document
# ---------------------------------------------------------------------------

def oMath(*items):
    om = _el(_m("oMath"), *items)
    return om


def oMathPara(om):
    """Wraps a single <m:oMath> into a display <m:oMathPara> (its own
    centered block, as Word does for equations on their own line)."""
    ompara = _el(_m("oMathPara"))
    ompr = _el(_m("oMathParaPr"))
    ompara.append(ompr)
    ompara.append(om)
    return ompara


def add_inline_math(paragraph, om):
    """Append an inline <m:oMath> to the end of a paragraph's XML."""
    paragraph._p.append(om)
    return paragraph


def add_display_math(paragraph, om):
    """Append a display <m:oMathPara> to the end of a paragraph's XML."""
    ompara = oMathPara(om)
    paragraph._p.append(ompara)
    return paragraph


def new_math_paragraph(doc_or_cell, om, style=None, alignment=None):
    """Create a brand-new paragraph containing one display math equation."""
    p = doc_or_cell.add_paragraph(style=style)
    if alignment is not None:
        p.alignment = alignment
    add_display_math(p, om)
    return p



# ---------------------------------------------------------------------------
# Convenience shortcuts for building common physics-equation fragments.
# Each shortcut returns a *list* of OMML nodes so callers can freely
# concatenate fragments with `+` before splatting them into oMath(*parts).
# ---------------------------------------------------------------------------

def V(letter, subscript=None, sub_italic=False):
    """An italic variable, e.g. V('v') -> v ; V('v','o') -> v_o ; V('x','o')-> x_o"""
    base = [mrun(letter)]
    if subscript is None:
        return base
    subitems = [mrun(str(subscript), italic=(None if sub_italic else False))]
    return [sub(base, subitems)]


def N(text):
    """A literal number / operator / unit chunk, rendered upright (non-italic)."""
    return [mrun(text, italic=False)]


def PW(base_items, exponent):
    """Superscript / power, e.g. PW(V('v'), '2') -> v^2"""
    return [sup(base_items, [mrun(str(exponent), italic=False)])]


def FRAC(num_items, den_items):
    return [frac(num_items, den_items)]


def PAREN(items, open_ch="(", close_ch=")"):
    return [paren(items, open_ch, close_ch)]


def SQRT(items):
    return [sqrt(items)]


def BOX(items):
    return [boxed(items)]


def cat(*groups):
    """Flatten/concatenate several fragment-lists into one list.

    IMPORTANT: lxml elements can only live in one place in the tree at a
    time -- appending an already-inserted element *moves* it rather than
    copying it. Since this module encourages building a fragment once
    (e.g. `v2 = SUBX(V('v'), '2')`) and reusing it across multiple
    equations/derivation rows, cat() deep-copies every element it receives
    so each usage gets its own independent copy.
    """
    import copy
    out = []
    for g in groups:
        for item in g:
            out.append(copy.deepcopy(item))
    return out


# ---------------------------------------------------------------------------
# Equals-sign aligned derivation block.
#
# Word/OOXML's official eqArr alignment-marker mechanism is inconsistently
# honored across renderers, so for guaranteed pixel-level alignment we use
# the well-established technique of a borderless 3-column table:
#   [ right-aligned LHS ] [ centered "=" ] [ left-aligned RHS ]
# Every row's "=" therefore lands in the same column, aligning all equal
# signs in the block -- while every cell's contents are still real OMML
# math (not plain text).
# ---------------------------------------------------------------------------

def add_aligned_derivation(doc_or_cell, rows, annotations=None, col_widths=None):
    """
    rows: list of (lhs_items, rhs_items) OMML-fragment-list pairs.
          Pass lhs_items=None to omit the LHS on a row (continuation line).
    annotations: optional list (same length as rows) of strings/None to show
                 as a small italic note to the right of that row.
    Returns the created table.
    """
    from docx.shared import Pt, Inches
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn as _qn

    n = len(rows)
    has_annot = annotations is not None
    ncols = 3 + (1 if has_annot else 0)
    table = doc_or_cell.add_table(rows=n, cols=ncols)
    table.autofit = True
    # Remove all borders (invisible table)
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:%s" % edge)
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)

    for i, row in enumerate(rows):
        lhs, rhs = row
        cells = table.rows[i].cells
        # LHS cell (right aligned)
        p_l = cells[0].paragraphs[0]
        p_l.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if lhs is not None:
            add_inline_math(p_l, oMath(*lhs))
        # "=" cell (centered)
        p_m = cells[1].paragraphs[0]
        p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if lhs is not None:
            add_inline_math(p_m, oMath(*N("=")))
        # RHS cell (left aligned)
        p_r = cells[2].paragraphs[0]
        p_r.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline_math(p_r, oMath(*rhs))
        # Annotation cell
        if has_annot:
            p_a = cells[3].paragraphs[0]
            p_a.alignment = WD_ALIGN_PARAGRAPH.LEFT
            note = annotations[i] if i < len(annotations) else None
            if note:
                from style_helpers import add_rich_text, parse_var_text, GRAY_TEXT
                add_rich_text(p_a, parse_var_text(note), italic=True, size=9.5,
                              color=GRAY_TEXT)

    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = w
    return table



def SUBX(base_items, subscript_text, sub_italic=False):
    """Subscript an arbitrary OMML fragment list, e.g. SUBX(PAREN(...), '1')
    -> (...)_1 ; or SUBX(V('v'), '1') -> v_1."""
    subitems = [mrun(str(subscript_text), italic=(None if sub_italic else False))]
    return [sub(base_items, subitems)]


def UPW(unit_text, exponent):
    """Upright (non-italic) base raised to a power, e.g. UPW('s','2') -> s^2,
    used for units like m/s^2 rather than for variables."""
    return [sup([mrun(unit_text, italic=False)], [mrun(str(exponent), italic=False)])]
