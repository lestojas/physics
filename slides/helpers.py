"""
Reusable design helpers for building a modern 16:9 PowerPoint deck.
Design system: colours, fonts, cards, bullets, tables, flow/funnel/chain diagrams.
"""
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- palette
INK        = "14213D"   # deep navy (primary text / dark bg)
INK_SOFT   = "24344F"
PRIMARY    = "1D7A8C"   # teal
PRIMARY_DK = "115968"
PRIMARY_LT = "E4F1F3"
ACCENT     = "F4795B"   # coral
ACCENT_LT  = "FCE4DD"
AMBER      = "F2B705"
AMBER_LT   = "FCF0CC"
SUCCESS    = "2A9D8F"
SUCCESS_LT = "D9F0EC"
DANGER     = "D1495B"
DANGER_LT  = "F7DEE2"
BG         = "F6F9FB"
CARD       = "FFFFFF"
MUTED      = "5B6B7B"
LINE       = "D8E1E8"
WHITE      = "FFFFFF"

# ---------------------------------------------------------------- fonts
F_TITLE = "Segoe UI Semibold"
F_HEAD  = "Segoe UI Semibold"
F_BODY  = "Segoe UI"
F_LIGHT = "Segoe UI Light"

SW = 13.333
SH = 7.5


def C(hexstr):
    return RGBColor.from_string(hexstr)


# ---------------------------------------------------------------- primitives
def _no_shadow(shape):
    shape.shadow.inherit = False


def soft_shadow(shape, blur=10, dist=5, alpha=76, color="1B2A3A"):
    """Add a subtle drop shadow to a shape for depth."""
    spPr = shape._element.spPr
    for old in spPr.findall(qn('a:effectLst')):
        spPr.remove(old)
    eff = spPr.makeelement(qn('a:effectLst'), {})
    sh = spPr.makeelement(qn('a:outerShdw'), {
        'blurRad': str(Emu(Pt(blur))),
        'dist': str(Emu(Pt(dist))),
        'dir': '5400000',
        'rotWithShape': '0',
    })
    clr = spPr.makeelement(qn('a:srgbClr'), {'val': color})
    a = spPr.makeelement(qn('a:alpha'), {'val': str(int((100 - alpha) * 1000))})
    clr.append(a)
    sh.append(clr)
    eff.append(sh)
    spPr.append(eff)


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0,
         kind=MSO_SHAPE.RECTANGLE, radius=None, shadow=False):
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    _no_shadow(sp)
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = C(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = C(line)
        sp.line.width = Pt(line_w)
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if shadow:
        soft_shadow(sp)
    sp.text_frame.word_wrap = True
    return sp


def rrect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.09, shadow=False):
    return rect(slide, x, y, w, h, fill=fill, line=line, line_w=line_w,
                kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius, shadow=shadow)


def _set_bullet(p, color):
    """Give a paragraph a coloured square bullet with a hanging indent."""
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', '274638')
    pPr.set('indent', '-274638')
    buClr = pPr.makeelement(qn('a:buClr'), {})
    srgb = pPr.makeelement(qn('a:srgbClr'), {'val': color})
    buClr.append(srgb)
    buSz = pPr.makeelement(qn('a:buSzPct'), {'val': '80000'})
    buFont = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'})
    buChar = pPr.makeelement(qn('a:buChar'), {'char': '\u25AA'})
    for el in (buClr, buSz, buFont, buChar):
        pPr.append(el)


def _run(p, text, size, color, bold=False, italic=False, font=F_BODY):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = C(color)
    return r


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf


def para(tf, first=False, align=PP_ALIGN.LEFT, space_before=0, space_after=6,
         line_spacing=1.05, bullet_color=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    if space_before is not None:
        p.space_before = Pt(space_before)
    if space_after is not None:
        p.space_after = Pt(space_after)
    if line_spacing is not None:
        p.line_spacing = line_spacing
    if bullet_color:
        _set_bullet(p, bullet_color)
    return p


def fill_frame(shape, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE,
               space_after=0, line_spacing=1.0):
    """runs: list of (text,size,color,bold,italic,font) tuples -> single paragraph."""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(3)
    tf.margin_bottom = Pt(3)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    p.space_after = Pt(space_after)
    for spec in runs:
        text, size, color = spec[0], spec[1], spec[2]
        bold = spec[3] if len(spec) > 3 else False
        italic = spec[4] if len(spec) > 4 else False
        font = spec[5] if len(spec) > 5 else F_BODY
        _run(p, text, size, color, bold, italic, font)
    return tf


# ---------------------------------------------------------------- backgrounds / chrome
LECTURE_FOOTER = "Lecture 3  \u00b7  Review of Literature for Identifying Research Problems"


def background(slide, color=BG):
    rect(slide, -0.1, -0.1, SW + 0.2, SH + 0.2, fill=color)


def deco_corner(slide, color=PRIMARY_LT):
    """Subtle geometric accent in the top-right corner."""
    c = rect(slide, SW - 2.2, -1.4, 3.4, 3.4, fill=color,
             kind=MSO_SHAPE.OVAL)
    c = rect(slide, SW - 0.9, 0.7, 1.4, 1.4, fill=WHITE, kind=MSO_SHAPE.OVAL)


def header(slide, eyebrow, title, page, title_size=30):
    """Standard content-slide header. Returns y where content may begin."""
    background(slide)
    # left accent bar
    rrect(slide, 0.55, 0.62, 0.14, 0.9, fill=PRIMARY, radius=0.5)
    # eyebrow
    _, tf = textbox(slide, 0.85, 0.55, 11.6, 0.32)
    p = para(tf, first=True, space_after=0)
    _run(p, eyebrow.upper(), 12.5, PRIMARY, bold=True, font=F_HEAD)
    # title
    _, tf = textbox(slide, 0.82, 0.85, 11.7, 0.9)
    p = para(tf, first=True, space_after=0, line_spacing=1.0)
    _run(p, title, title_size, INK, bold=True, font=F_TITLE)
    # accent rule
    rrect(slide, 0.85, 1.62, 0.75, 0.07, fill=ACCENT, radius=0.5)
    footer(slide, page)
    return 1.9


def footer(slide, page):
    _, tf = textbox(slide, 0.85, 7.02, 9.0, 0.3)
    p = para(tf, first=True, space_after=0)
    _run(p, LECTURE_FOOTER, 9.5, MUTED, font=F_BODY)
    # page badge
    b = rect(slide, 12.42, 6.92, 0.5, 0.5, fill=INK, kind=MSO_SHAPE.OVAL)
    fill_frame(b, [(str(page), 12, WHITE, True)], align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------- content components
def bullets(slide, x, y, w, items, size=24, color=INK, bullet_color=PRIMARY,
            gap=9, line_spacing=1.05, h=4.8):
    """items: list of strings OR list of run-lists [(text,bold,italic,color?)]."""
    _, tf = textbox(slide, x, y, w, h)
    first = True
    for it in items:
        p = para(tf, first=first, space_after=gap, line_spacing=line_spacing,
                 bullet_color=bullet_color)
        first = False
        if isinstance(it, str):
            _run(p, it, size, color)
        else:
            for seg in it:
                t = seg[0]
                bold = seg[1] if len(seg) > 1 else False
                italic = seg[2] if len(seg) > 2 else False
                col = seg[3] if len(seg) > 3 else color
                sz = seg[4] if len(seg) > 4 else size
                _run(p, t, sz, col, bold, italic)
    return tf


def checklist(slide, x, y, w, items, size=24, color=INK, mark_color=SUCCESS, gap=0.16, row_h=0.62):
    cy = y
    for it in items:
        badge = rect(slide, x, cy, 0.42, 0.42, fill=mark_color, kind=MSO_SHAPE.OVAL)
        fill_frame(badge, [("\u2713", 15, WHITE, True)], align=PP_ALIGN.CENTER)
        _, tf = textbox(slide, x + 0.62, cy - 0.06, w - 0.62, row_h, anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, first=True, space_after=0, line_spacing=1.0)
        if isinstance(it, str):
            _run(p, it, size, color)
        else:
            for seg in it:
                t = seg[0]
                bold = seg[1] if len(seg) > 1 else False
                italic = seg[2] if len(seg) > 2 else False
                col = seg[3] if len(seg) > 3 else color
                _run(p, t, size, col, bold, italic)
        cy += row_h + gap
    return cy


def card(slide, x, y, w, h, fill=CARD, line=LINE, shadow=True, radius=0.06):
    sp = rrect(slide, x, y, w, h, fill=fill, line=line, line_w=1.0, radius=radius, shadow=shadow)
    return sp


def card_header(slide, x, y, w, title, band_color, h=0.62, text_color=WHITE, size=18,
                icon=None):
    band = rrect(slide, x, y, w, h, fill=band_color, radius=0.12)
    tx = x + 0.28
    if icon:
        ic = rect(slide, x + 0.24, y + (h - 0.4) / 2, 0.4, 0.4, fill=WHITE, kind=MSO_SHAPE.OVAL)
        fill_frame(ic, [(icon, 15, band_color, True)], align=PP_ALIGN.CENTER)
        tx = x + 0.82
    _, tf = textbox(slide, tx, y, w - (tx - x) - 0.2, h, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tf, first=True, space_after=0)
    _run(p, title, size, text_color, bold=True, font=F_HEAD)
    return band


def two_column(slide, top, left, right, bottom=6.65, gap=0.5, x0=0.85):
    """left/right: dict(title, band, items(list), item_color, size)."""
    total = SW - 2 * x0
    cw = (total - gap) / 2
    for col, cx in ((left, x0), (right, x0 + cw + gap)):
        h = bottom - top
        card(slide, cx, top, cw, h)
        card_header(slide, cx, top, cw, col["title"], col["band"],
                    icon=col.get("icon"), size=col.get("head_size", 18))
        items = col["items"]
        bullets(slide, cx + 0.32, top + 0.85, cw - 0.6, items,
                size=col.get("size", 22), bullet_color=col["band"],
                color=INK, gap=col.get("gap", 8), h=h - 1.0)
    return cw


def styled_table(slide, x, y, w, data, col_widths=None, header=True,
                 header_fill=PRIMARY, header_color=WHITE, size=17, header_size=17,
                 row_h=0.5, header_h=0.55, cell_colors=None, align_first_left=True,
                 zebra=("FFFFFF", "EEF4F7")):
    """data: list of rows, each a list of cell strings. Returns the table shape."""
    rows = len(data)
    cols = len(data[0])
    total_h = header_h + (rows - 1) * row_h if header else rows * row_h
    gtbl = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(total_h))
    tbl = gtbl.table
    tbl.first_row = False
    tbl.horz_banding = False
    # column widths
    if col_widths:
        tot = sum(col_widths)
        for i, cwd in enumerate(col_widths):
            tbl.columns[i].width = Inches(w * cwd / tot)
    # row heights
    for ri in range(rows):
        tbl.rows[ri].height = Inches(header_h if (header and ri == 0) else row_h)
    for ri in range(rows):
        for ci in range(cols):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Pt(8)
            cell.margin_right = Pt(8)
            cell.margin_top = Pt(4)
            cell.margin_bottom = Pt(4)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            is_head = header and ri == 0
            if is_head:
                cell.fill.solid(); cell.fill.fore_color.rgb = C(header_fill)
            elif cell_colors and (ri, ci) in cell_colors:
                cell.fill.solid(); cell.fill.fore_color.rgb = C(cell_colors[(ri, ci)])
            else:
                idx = ri if not header else ri - 1
                cell.fill.solid()
                cell.fill.fore_color.rgb = C(zebra[idx % 2])
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.0
            if align_first_left:
                p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.LEFT
            txt = data[ri][ci]
            r = p.add_run(); r.text = txt
            r.font.size = Pt(header_size if is_head else size)
            r.font.name = F_HEAD if is_head else F_BODY
            r.font.bold = True if is_head else False
            r.font.color.rgb = C(header_color if is_head else INK)
    return gtbl


def down_arrow(slide, cx, y, color=PRIMARY, w=0.34, h=0.28):
    a = rect(slide, cx - w / 2, y, w, h, fill=color, kind=MSO_SHAPE.DOWN_ARROW)
    return a


def right_arrow(slide, x, cy, color=PRIMARY, w=0.34, h=0.32):
    a = rect(slide, x, cy - h / 2, w, h, fill=color, kind=MSO_SHAPE.RIGHT_ARROW)
    return a


def vchain(slide, x, y, w, steps, box_h=0.66, gap=0.24, colors=None,
           text_size=17, box_fill=None, arrow_color=PRIMARY):
    """steps: list of (label, body) OR strings. Vertical connected boxes."""
    cy = y
    n = len(steps)
    for i, st in enumerate(steps):
        col = (colors[i] if colors else PRIMARY)
        b = rrect(slide, x, cy, w, box_h, fill=(box_fill or col), radius=0.12, shadow=True)
        if isinstance(st, tuple):
            label, body = st
            tf = b.text_frame
            tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Pt(10); tf.margin_right = Pt(10)
            p = tf.paragraphs[0]; p.line_spacing = 1.0
            _run(p, label + "  ", text_size, WHITE, bold=True, font=F_HEAD)
            if body:
                _run(p, body, text_size - 1, "EAF6F8", italic=False)
        else:
            fill_frame(b, [(st, text_size, WHITE, True)], align=PP_ALIGN.LEFT)
        cy += box_h
        if i < n - 1:
            down_arrow(slide, x + w / 2, cy + gap * 0.12, color=arrow_color,
                       h=gap * 0.76)
            cy += gap
    return cy
